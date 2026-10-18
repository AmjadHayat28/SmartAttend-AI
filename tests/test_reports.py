"""Teacher attendance reports + Excel export: own active classes only, finalized sessions only."""

from datetime import date, timedelta
from io import BytesIO

import pytest
from openpyxl import load_workbook

from backend.app.core.config import settings
from backend.app.services.attendance_policy import summarize
from tests.conftest import STUDENT_EMAIL
from tests.test_attendance import API, cls, finalize, mark, start, teacher_h  # noqa: F401

TODAY = date.today()
PREV = TODAY.replace(day=1) - timedelta(days=1)  # last day of the previous month
XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

# CS-321 / A roster (cls fixture): F23CSE01, 03 (pending), 04, 05, 06.
PLAN = [
    (PREV, 1, {"F23CSE01": "present", "F23CSE04": "absent"}),
    (TODAY, 1, {"F23CSE01": "absent", "F23CSE03": "late", "F23CSE04": "absent", "F23CSE06": "excused"}),
    (TODAY, 2, {"F23CSE01": "late", "F23CSE04": "absent"}),
]


def _session(client, teacher_h, assignment, day, number, statuses, final=True):
    s = start(client, teacher_h, assignment, day=day.isoformat(), number=number).json()["session"]
    s = mark(client, teacher_h, s, [statuses.get(r["registration_no"], "present") for r in s["records"]]).json()
    if final:
        res = finalize(client, teacher_h, s)
        assert res.status_code == 200, res.text
        return res.json()
    return s


@pytest.fixture()
def data(client, teacher_h, cls, monkeypatch):
    """Three finalized sessions (two today, one last month) plus a deleted, a draft and a cancelled
    session — all marking everyone absent, so any leak would show up in the counts."""
    monkeypatch.setattr(settings, "ATTENDANCE_MAX_BACKDATE_DAYS", 45)
    for day, number, statuses in PLAN:
        _session(client, teacher_h, cls["assignment"], day, number, statuses)
    everyone_absent = {r: "absent" for r in ("F23CSE01", "F23CSE03", "F23CSE04", "F23CSE05", "F23CSE06")}
    deleted = _session(client, teacher_h, cls["assignment"], TODAY, 3, everyone_absent)
    assert client.post(f"{API}/attendance/sessions/{deleted['id']}/delete", headers=teacher_h,
                       json={"reason": "Taken by mistake"}).status_code == 204
    _session(client, teacher_h, cls["assignment"], TODAY, 4, everyone_absent, final=False)  # draft
    cancelled = _session(client, teacher_h, cls["assignment"], TODAY, 5, everyone_absent, final=False)
    assert client.post(f"{API}/attendance/sessions/{cancelled['id']}/cancel", headers=teacher_h,
                       json={"version": cancelled["version"], "reason": "Wrong slot"}).status_code == 200
    return cls


def _expected(reg, sessions=PLAN):
    statuses = [plan.get(reg, "present") for _, _, plan in sessions]
    return summarize({st: statuses.count(st) for st in set(statuses)})


def _row(report, reg):
    return next(s for s in report["students"] if s["registration_no"] == reg)


# ── Authorization ─────────────────────────────────────────────────────────


def test_teacher_sees_only_own_active_classes(client, teacher_h, data):
    classes = client.get(f"{API}/reports/classes", headers=teacher_h).json()
    assert {(c["course_code"], c["section_name"]) for c in classes} == {("CS-321", "A"), ("CS-321L", "A")}
    assert client.get(f"{API}/reports/class/{data['assignment']}", headers=teacher_h).status_code == 200


def _all_endpoints(assignment, student_id):
    return [
        ("/reports/classes", {}),
        (f"/reports/class/{assignment}", {}),
        (f"/reports/class/{assignment}/daily", {"date": TODAY.isoformat()}),
        (f"/reports/class/{assignment}/monthly", {"year": TODAY.year, "month": TODAY.month}),
        (f"/reports/class/{assignment}/students/{student_id}", {}),
        (f"/reports/class/{assignment}/export", {"type": "summary"}),
        (f"/reports/class/{assignment}/export", {"type": "daily", "date": TODAY.isoformat()}),
        (f"/reports/class/{assignment}/export", {"type": "monthly", "year": TODAY.year, "month": TODAY.month}),
        (f"/reports/class/{assignment}/export", {"type": "student", "student_id": student_id}),
    ]


@pytest.mark.parametrize("who", ["hod", "admin", "student"])
def test_non_teachers_get_403(client, who, hod_h, admin_h, login, data, ids):
    headers = {"hod": hod_h, "admin": admin_h}.get(who) or login(STUDENT_EMAIL)
    for path, params in _all_endpoints(data["assignment"], ids["student"]):
        assert client.get(f"{API}{path}", params=params, headers=headers).status_code == 403, (who, path)


def test_other_teachers_assignment_is_404(client, hod_h, ee_hod_h, ee, login, data, ids):
    # Same department, assigned to another class of the same section.
    client.post(f"{API}/teachers", headers=hod_h, json={
        "full_name": "Co Teacher", "email": "co@test.edu", "password": "Teach@1234", "employee_id": "EMP-T-777"})
    co_id = client.get(f"{API}/teachers", params={"q": "EMP-T-777"}, headers=hod_h).json()["items"][0]["id"]
    client.post(f"{API}/assignments", headers=hod_h, json={
        "teacher_id": co_id, "course_id": ids["courses"]["MA-301"], "section_id": ids["section_a"]})
    # Another department's teacher.
    client.post(f"{API}/assignments", headers=ee_hod_h, json={
        "teacher_id": ee["teacher"], "course_id": ee["course"], "section_id": ee["section"]})
    for headers in (login("co@test.edu", "Teach@1234"), login("teacher.ee@test.edu")):
        assert len(client.get(f"{API}/reports/classes", headers=headers).json()) == 1
        for path, params in _all_endpoints(data["assignment"], ids["student"])[1:]:
            assert client.get(f"{API}{path}", params=params, headers=headers).status_code == 404, path


def test_removed_assignment_is_no_longer_reportable(client, teacher_h, hod_h, data):
    assert client.delete(f"{API}/assignments/{data['assignment']}", headers=hod_h).status_code == 204
    assert client.get(f"{API}/reports/class/{data['assignment']}", headers=teacher_h).status_code == 404


# ── Calculation ───────────────────────────────────────────────────────────


def test_class_summary_counts_only_finalized(client, teacher_h, data):
    report = client.get(f"{API}/reports/class/{data['assignment']}", headers=teacher_h).json()
    assert report["finalized_sessions"] == 3  # not the deleted, draft or cancelled ones
    assert report["enrolled_students"] == 5 and report["threshold"] == settings.LOW_ATTENDANCE_THRESHOLD
    assert report["report_class"]["course_code"] == "CS-321" and report["report_class"]["teacher_name"] == "Engr. Bilal Ahmed"
    assert [s["registration_no"] for s in report["students"]] == ["F23CSE01", "F23CSE03", "F23CSE04", "F23CSE05", "F23CSE06"]
    for reg in ("F23CSE01", "F23CSE03", "F23CSE04", "F23CSE05", "F23CSE06"):
        row, expected = _row(report, reg), _expected(reg)
        assert (row["present"], row["absent"], row["late"], row["excused"]) == (
            expected.present, expected.absent, expected.late, expected.excused), reg
        assert row["total_classes"] == 3 and row["percentage"] == expected.percentage, reg
        assert row["low_attendance"] == expected.is_low, reg
    one = _row(report, "F23CSE01")
    assert (one["present"], one["absent"], one["late"], one["excused"]) == (1, 1, 1, 0)
    assert (_row(report, "F23CSE06")["excused"], _row(report, "F23CSE03")["late"]) == (1, 1)
    assert _row(report, "F23CSE04")["percentage"] == 0.0 and _row(report, "F23CSE04")["low_attendance"] is True
    all_statuses = {}
    for reg in ("F23CSE01", "F23CSE03", "F23CSE04", "F23CSE05", "F23CSE06"):
        for _, _, plan in PLAN:
            st = plan.get(reg, "present")
            all_statuses[st] = all_statuses.get(st, 0) + 1
    assert report["class_percentage"] == summarize(all_statuses).percentage
    assert report["low_attendance_students"] == sum(1 for s in report["students"] if s["low_attendance"])


def test_low_attendance_uses_the_one_configured_threshold(client, teacher_h, data, monkeypatch):
    pct = _expected("F23CSE01").percentage  # 66.7 under the default policy
    monkeypatch.setattr(settings, "LOW_ATTENDANCE_THRESHOLD", pct - 1)
    report = client.get(f"{API}/reports/class/{data['assignment']}", headers=teacher_h).json()
    assert report["threshold"] == pct - 1 and _row(report, "F23CSE01")["low_attendance"] is False
    monkeypatch.setattr(settings, "LOW_ATTENDANCE_THRESHOLD", pct + 1)
    report = client.get(f"{API}/reports/class/{data['assignment']}", headers=teacher_h).json()
    assert _row(report, "F23CSE01")["low_attendance"] is True


def test_policy_changes_flow_through(client, teacher_h, data, monkeypatch):
    monkeypatch.setattr(settings, "ATTENDANCE_LATE_WEIGHT", 0.5)
    monkeypatch.setattr(settings, "ATTENDANCE_EXCUSED_POLICY", "absent")
    report = client.get(f"{API}/reports/class/{data['assignment']}", headers=teacher_h).json()
    for reg in ("F23CSE01", "F23CSE03", "F23CSE06"):
        assert _row(report, reg)["percentage"] == _expected(reg).percentage, reg


def test_zero_finalized_is_na(client, teacher_h, cls):
    report = client.get(f"{API}/reports/class/{cls['assignment']}", headers=teacher_h).json()
    assert report["finalized_sessions"] == 0 and report["class_percentage"] is None
    assert all(s["percentage"] is None and not s["low_attendance"] for s in report["students"])


# ── Filters ───────────────────────────────────────────────────────────────


def test_daily_report(client, teacher_h, data):
    report = client.get(f"{API}/reports/class/{data['assignment']}/daily", params={"date": TODAY.isoformat()},
                        headers=teacher_h).json()
    assert [s["session_number"] for s in report["sessions"]] == [1, 2]  # no deleted #3, draft #4, cancelled #5
    first = {r["registration_no"]: r["status"] for r in report["sessions"][0]["records"]}
    assert first == {"F23CSE01": "absent", "F23CSE03": "late", "F23CSE04": "absent", "F23CSE05": "present",
                     "F23CSE06": "excused"}
    second = {r["registration_no"]: r["status"] for r in report["sessions"][1]["records"]}
    assert second["F23CSE01"] == "late" and second["F23CSE05"] == "present"
    for record in report["sessions"][0]["records"]:
        assert set(record) == {"student_id", "registration_no", "student_name", "status"}
    searched = client.get(f"{API}/reports/class/{data['assignment']}/daily",
                          params={"date": TODAY.isoformat(), "q": "cse04"}, headers=teacher_h).json()
    assert [r["registration_no"] for s in searched["sessions"] for r in s["records"]] == ["F23CSE04", "F23CSE04"]
    empty_day = (TODAY - timedelta(days=60)).isoformat()
    assert client.get(f"{API}/reports/class/{data['assignment']}/daily", params={"date": empty_day},
                      headers=teacher_h).json()["sessions"] == []
    assert client.get(f"{API}/reports/class/{data['assignment']}/daily", headers=teacher_h).status_code == 422


def test_monthly_report_excludes_other_months(client, teacher_h, data):
    this = client.get(f"{API}/reports/class/{data['assignment']}/monthly",
                      params={"year": TODAY.year, "month": TODAY.month}, headers=teacher_h).json()
    assert this["finalized_sessions"] == 2 and this["date_from"] == TODAY.replace(day=1).isoformat()
    today_plan = [p for p in PLAN if p[0] == TODAY]
    one = _row(this, "F23CSE01")
    assert (one["total_classes"], one["absent"], one["late"], one["present"]) == (2, 1, 1, 0)
    assert one["percentage"] == _expected("F23CSE01", today_plan).percentage
    prev = client.get(f"{API}/reports/class/{data['assignment']}/monthly",
                      params={"year": PREV.year, "month": PREV.month}, headers=teacher_h).json()
    assert prev["finalized_sessions"] == 1 and _row(prev, "F23CSE01")["present"] == 1
    empty = client.get(f"{API}/reports/class/{data['assignment']}/monthly", params={"year": 2001, "month": 1},
                       headers=teacher_h).json()
    assert empty["finalized_sessions"] == 0 and empty["class_percentage"] is None
    assert client.get(f"{API}/reports/class/{data['assignment']}/monthly", params={"year": 2026, "month": 13},
                      headers=teacher_h).status_code == 422


def test_date_range_and_search(client, teacher_h, data):
    report = client.get(f"{API}/reports/class/{data['assignment']}", headers=teacher_h,
                        params={"date_from": TODAY.isoformat(), "date_to": TODAY.isoformat()}).json()
    assert report["finalized_sessions"] == 2 and _row(report, "F23CSE01")["total_classes"] == 2
    only_prev = client.get(f"{API}/reports/class/{data['assignment']}", headers=teacher_h,
                           params={"date_to": PREV.isoformat()}).json()
    assert only_prev["finalized_sessions"] == 1
    assert client.get(f"{API}/reports/class/{data['assignment']}", headers=teacher_h,
                      params={"date_from": TODAY.isoformat(), "date_to": PREV.isoformat()}).status_code == 422
    for q, regs in (("cse05", ["F23CSE05"]), ("raza", ["F23CSE01"]), ("nobody", [])):
        found = client.get(f"{API}/reports/class/{data['assignment']}", params={"q": q}, headers=teacher_h).json()
        assert [s["registration_no"] for s in found["students"]] == regs, q


def test_student_report(client, teacher_h, data, ids):
    res = client.get(f"{API}/reports/class/{data['assignment']}/students/{ids['student']}", headers=teacher_h)
    assert res.status_code == 200
    report = res.json()
    assert report["student"]["registration_no"] == "F23CSE01" and report["student"]["total_classes"] == 3
    assert report["student"]["percentage"] == _expected("F23CSE01").percentage
    assert [(r["session_date"], r["session_number"], r["status"]) for r in report["records"]] == [
        (TODAY.isoformat(), 2, "late"), (TODAY.isoformat(), 1, "absent"), (PREV.isoformat(), 1, "present")]
    ranged = client.get(f"{API}/reports/class/{data['assignment']}/students/{ids['student']}", headers=teacher_h,
                        params={"date_from": TODAY.isoformat()}).json()
    assert len(ranged["records"]) == 2 and ranged["student"]["total_classes"] == 2


def test_student_outside_class_is_404(client, teacher_h, hod_h, data, ids, ee):
    outsider = client.post(f"{API}/students", headers=hod_h, json={
        "full_name": "Other Section", "email": "b9@test.edu", "password": "Student@123", "registration_no": "F23CSE90",
        "batch_id": ids["batch"], "semester_id": ids["semester"], "section_id": ids["section_b"]}).json()
    for student_id in (outsider["id"], ee["student"], 999999):
        assert client.get(f"{API}/reports/class/{data['assignment']}/students/{student_id}",
                          headers=teacher_h).status_code == 404


def test_unenrolled_student_history_is_kept(client, teacher_h, data):
    enrolled = client.get(f"{API}/enrollments", params={"assignment_id": data["assignment"]}, headers=teacher_h).json()
    four = next(e for e in enrolled["items"] if e["registration_no"] == "F23CSE04")
    assert client.delete(f"{API}/enrollments/{four['id']}", headers=teacher_h).status_code == 204
    report = client.get(f"{API}/reports/class/{data['assignment']}", headers=teacher_h).json()
    row = _row(report, "F23CSE04")
    assert row["currently_enrolled"] is False and row["absent"] == 3
    assert report["enrolled_students"] == 4
    assert client.get(f"{API}/reports/class/{data['assignment']}/students/{four['student_id']}",
                      headers=teacher_h).status_code == 200


# ── Excel ─────────────────────────────────────────────────────────────────


def _workbook(res):
    assert res.status_code == 200, res.text
    assert res.headers["content-type"] == XLSX
    assert res.content[:2] == b"PK"  # a real zip-based .xlsx, not CSV
    return load_workbook(BytesIO(res.content))


def _table(ws, first_header):
    """(header row index, headers, data rows as dicts) for the table starting at `first_header`."""
    for row in ws.iter_rows():
        if row[0].value == first_header:
            headers = [c.value for c in row if c.value is not None]
            rows = []
            for data_row in ws.iter_rows(min_row=row[0].row + 1, values_only=True):
                if data_row[0] is None:
                    break
                rows.append(dict(zip(headers, data_row)))
            return row[0].row, headers, rows
    raise AssertionError(f"no table starting with {first_header!r}")


def test_summary_export(client, teacher_h, data):
    api = client.get(f"{API}/reports/class/{data['assignment']}", headers=teacher_h).json()
    res = client.get(f"{API}/reports/class/{data['assignment']}/export", params={"type": "summary"}, headers=teacher_h)
    wb = _workbook(res)
    assert 'filename="CS321_SectionA_Attendance_Summary.xlsx"' in res.headers["content-disposition"]
    ws = wb["Summary"]
    assert ws["A1"].value == "SmartClass AI Attendance Management System"
    assert ws["A2"].value == "Class Attendance Summary"
    labels = {ws.cell(row=r, column=1).value: ws.cell(row=r, column=2).value for r in range(4, 14)}
    assert labels["Course"] == "CS-321 — Digital System Design" and labels["Teacher"] == "Engr. Bilal Ahmed"
    assert "Section A" in labels["Section"] and labels["Generated"]
    header_row, headers, rows = _table(ws, "Registration Number")
    assert headers[:8] == ["Registration Number", "Student Name", "Total Classes", "Present", "Absent", "Late",
                           "Excused", "Attendance %"]
    assert ws.cell(row=header_row, column=1).font.bold and ws.freeze_panes == f"A{header_row + 1}"
    assert [r["Registration Number"] for r in rows] == [s["registration_no"] for s in api["students"]]
    for r, s in zip(rows, api["students"]):
        assert (r["Present"], r["Absent"], r["Late"], r["Excused"]) == (s["present"], s["absent"], s["late"], s["excused"])
        assert r["Attendance %"] == pytest.approx(s["percentage"] / 100)
        assert (r["Low Attendance"] == "Low Attendance") == s["low_attendance"]
    pct_cell = ws.cell(row=header_row + 1, column=8)
    assert pct_cell.number_format == "0.0%"
    text = " ".join(str(c.value) for row in ws.iter_rows() for c in row if c.value is not None).lower()
    assert "confidence" not in text and "upload" not in text and ".jpg" not in text


def test_daily_monthly_student_exports(client, teacher_h, data, ids):
    res = client.get(f"{API}/reports/class/{data['assignment']}/export",
                     params={"type": "daily", "date": TODAY.isoformat()}, headers=teacher_h)
    ws = _workbook(res)["Daily"]
    assert f'filename="CS321_SectionA_{TODAY.isoformat()}_Attendance.xlsx"' in res.headers["content-disposition"]
    assert ws["A2"].value == "Daily Attendance Report"
    _, headers, rows = _table(ws, "Registration Number")
    assert headers == ["Registration Number", "Student Name", "Date", "Session Number", "Status"]
    assert len(rows) == 10 and {r["Session Number"] for r in rows} == {1, 2}
    first = {r["Registration Number"]: r["Status"] for r in rows if r["Session Number"] == 1}
    assert first["F23CSE03"] == "Late" and first["F23CSE06"] == "Excused" and first["F23CSE01"] == "Absent"

    month = f"{TODAY.year:04d}-{TODAY.month:02d}"
    api = client.get(f"{API}/reports/class/{data['assignment']}/monthly",
                     params={"year": TODAY.year, "month": TODAY.month}, headers=teacher_h).json()
    res = client.get(f"{API}/reports/class/{data['assignment']}/export",
                     params={"type": "monthly", "year": TODAY.year, "month": TODAY.month}, headers=teacher_h)
    ws = _workbook(res)["Monthly"]
    assert f'filename="CS321_SectionA_{month}_Attendance.xlsx"' in res.headers["content-disposition"]
    _, _, rows = _table(ws, "Registration Number")
    one = next(r for r in rows if r["Registration Number"] == "F23CSE01")
    assert one["Total Classes"] == 2 and one["Attendance %"] == pytest.approx(_row(api, "F23CSE01")["percentage"] / 100)

    res = client.get(f"{API}/reports/class/{data['assignment']}/export",
                     params={"type": "student", "student_id": ids["student"]}, headers=teacher_h)
    ws = _workbook(res)["Student"]
    assert 'filename="CS321_SectionA_F23CSE01_Attendance.xlsx"' in res.headers["content-disposition"]
    _, _, summary = _table(ws, "Registration Number")
    assert summary[0]["Total Classes"] == 3
    _, headers, history = _table(ws, "Date")
    assert headers == ["Date", "Session Number", "Status"]
    assert [(r["Session Number"], r["Status"]) for r in history] == [(2, "Late"), (1, "Absent"), (1, "Present")]


def test_export_validation(client, teacher_h, data):
    base = f"{API}/reports/class/{data['assignment']}/export"
    assert client.get(base, params={"type": "daily"}, headers=teacher_h).status_code == 422
    assert client.get(base, params={"type": "monthly", "year": 2026}, headers=teacher_h).status_code == 422
    assert client.get(base, params={"type": "student"}, headers=teacher_h).status_code == 422
    assert client.get(base, params={"type": "pdf"}, headers=teacher_h).status_code == 422
